"""Turn the app's canonical, already-escaped print pages into editable Word blocks.

This handles the contract templates' small HTML vocabulary, not arbitrary HTML/CSS.
Only embedded PNGs are accepted: export never retrieves a URL or local file.
"""
import base64
import re
from io import BytesIO

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


def contract_docx(html, *, blank=False):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(14)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, 18), (2, 12), (3, 10)):
        style = document.styles[f'Heading {level}']
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor.from_string('222222')
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    properties = document.core_properties
    properties.title = 'Petrol Zone Sales Contract' + (' — Blank Template' if blank else '')
    properties.author = 'Petrol Zone'
    properties.last_modified_by = 'Petrol Zone'
    soup = BeautifulSoup(html, 'html.parser')
    # Running markers also cover overflow pages in long product tables/clauses.
    header = soup.select_one('#header-html .pz-running')
    footer = soup.select_one('#footer-html .pz-footer')
    for paragraph, value in (
            (section.header.paragraphs[0], header.get_text() if header else 'PETROL ZONE · BLANK TEMPLATE'),
            (section.footer.paragraphs[0], footer.get_text() if footer else 'PETROL ZONE')):
        paragraph.add_run(_text(value)).font.size = Pt(7)
        paragraph.paragraph_format.space_after = Pt(0)
    section.header_distance = section.footer_distance = Mm(6)
    # Select only printable pages, excluding HTML toolbar and PDF-only wrappers.
    for index, page in enumerate(soup.select('section.page')):
        if index:
            document.add_page_break()
        for child in page.children:
            _block(child, document, Mm(182))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _text(value):
    # Word's XML cannot represent these control characters. Keep tabs/newlines,
    # Unicode and literal markup; python-docx handles XML escaping itself.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]', '', str(value))


def _inline(node, paragraph, *, bold=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(_text(node))
        if bold:
            run.bold = True
        return
    if node.name in ('script', 'style', 'button'):
        return
    if node.name == 'br':
        paragraph.add_run().add_break()
    elif node.name == 'img':
        source = node.get('src', '')
        if source.startswith('data:image/png;base64,'):
            data = base64.b64decode(source.partition(',')[2], validate=True)
            width = Mm(35 if 'seller-stamp' in node.get('class', []) else 19)
            paragraph.add_run().add_picture(BytesIO(data), width=width)
    elif 'line' in node.get('class', []):
        paragraph.add_run('____________________')
    else:
        for child in node.children:
            _inline(child, paragraph, bold=bold or node.name in ('strong', 'b'))


def _shade(cell, color):
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(element)


def _masthead(node, parent, width):
    table = parent.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Mm(24), width - Mm(70), Mm(46)]
    for column, cell, size in zip(table.columns, table.rows[0].cells, widths):
        column.width = cell.width = size
    cells = table.rows[0].cells
    image = node.find('img')
    if image:
        _inline(image, cells[0].paragraphs[0])
    brand = cells[1].paragraphs[0]
    run = brand.add_run('PETROL ZONE')
    run.bold, run.font.size = True, Pt(18)
    subtitle = node.find('span')
    if subtitle:
        run = brand.add_run('\n' + subtitle.get_text())
        run.font.size = Pt(7)
    reference = node.find(class_='reference')
    if reference:
        paragraph = cells[2].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _inline(reference, paragraph)
    borders = OxmlElement('w:tblBorders')
    bottom = OxmlElement('w:bottom')
    for key, value in dict(val='single', sz='16', color='EFA90A').items():
        bottom.set(qn('w:' + key), value)
    borders.append(bottom)
    table._tbl.tblPr.append(borders)
    paragraph = parent.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = Pt(2)


def _table(node, parent, width):
    rows = [row for row in node.find_all('tr') if row.find_parent('table') is node]
    if not rows:
        return
    columns = len(rows[0].find_all(['td', 'th'], recursive=False))
    table = parent.add_table(rows=0, cols=columns)
    table.autofit = False
    table.style = 'Table Grid'
    margins = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        margin = OxmlElement('w:' + side)
        margin.set(qn('w:w'), '60')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    table._tbl.tblPr.append(margins)
    first_cells = rows[0].find_all(['td', 'th'], recursive=False)
    widths = []
    for cell in first_cells:
        match = re.search(r'width:\s*([\d.]+)%', cell.get('style', ''))
        widths.append(int(width * float(match[1]) / 100) if match else int(width / columns))
    # Partial width declarations (blank product table) leave the rest equally split.
    specified = [i for i, cell in enumerate(first_cells) if 'width:' in cell.get('style', '')]
    if specified and len(specified) < columns:
        remaining = int((width - sum(widths[i] for i in specified)) / (columns - len(specified)))
        widths = [size if i in specified else remaining for i, size in enumerate(widths)]
    for column, size in zip(table.columns, widths):
        column.width = size
    classes = node.get('class', [])
    for html_row in rows:
        row = table.add_row()
        cells = html_row.find_all(['td', 'th'], recursive=False)
        header = any(cell.name == 'th' for cell in cells)
        if header:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        if 'products' in classes and not header:
            row.height, row.height_rule = Mm(24), WD_ROW_HEIGHT_RULE.AT_LEAST
        if 'signatures' in classes or 'keep' in classes:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, html_cell, size in zip(row.cells, cells, widths):
            cell.width = size
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            for child in html_cell.children:
                if getattr(child, 'name', None) in ('table', 'div'):
                    _block(child, cell, size - Mm(4))
                    paragraph = cell.add_paragraph()
                else:
                    _inline(child, paragraph, bold=header)
                    if getattr(child, 'name', None) == 'strong' and 'details' in classes:
                        paragraph.add_run().add_break()
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    if header:
                        run.font.color.rgb = RGBColor.from_string('FFC044')
            if header:
                _shade(cell, '262626')
    parent.add_paragraph().paragraph_format.space_after = Pt(1)


def _block(node, parent, width):
    if isinstance(node, NavigableString):
        if node.strip():
            parent.add_paragraph(_text(node))
        return
    classes = node.get('class', [])
    if 'masthead' in classes:
        _masthead(node, parent, width)
    elif node.name == 'table':
        _table(node, parent, width)
    elif node.name == 'img':
        _inline(node, parent.add_paragraph())
    elif 'ruled' in classes or 'signature' in classes:
        paragraph = parent.add_paragraph('_' * min(64, max(10, int(width / Mm(2)))))
        paragraph.paragraph_format.space_before = Pt(10)
    elif node.name in ('p', 'h1', 'h2', 'h3') or any(
            name in classes for name in ('draft', 'status', 'product-description')):
        paragraph = (parent.add_heading(level=int(node.name[1]))
            if node.name in ('h1', 'h2', 'h3') else parent.add_paragraph())
        _inline(node, paragraph)
        if 'draft' in classes:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold, run.font.size = True, Pt(15)
                run.font.color.rgb = RGBColor.from_string('AA3535' if 'cancelled' in classes else '815200')
        if node.find_parent(class_='terms'):
            paragraph.paragraph_format.keep_together = True
    else:
        for child in node.children:
            _block(child, parent, width)
