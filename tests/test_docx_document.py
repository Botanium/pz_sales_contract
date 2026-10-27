import base64
from io import BytesIO
from pathlib import Path
import unittest
from zipfile import ZipFile

from docx import Document
from lxml import etree

from pz_sales_contract.docx_document import contract_docx

ROOT = Path(__file__).resolve().parents[1]
PNG = (ROOT / 'pz_sales_contract/public/petrol_zone_logo.png').read_bytes()
URI = 'data:image/png;base64,' + base64.b64encode(PNG).decode()
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


class WordDocumentTests(unittest.TestCase):
    def test_bold_term_bullets_are_editable_word_list_items(self):
        data = contract_docx('<section class="page terms"><ul><li><strong>Important term</strong></li></ul></section>')
        paragraph = next(p for p in Document(BytesIO(data)).paragraphs if 'Important term' in p.text)
        self.assertEqual(paragraph.style.name, 'List Bullet')
        self.assertTrue(paragraph.runs[0].bold)

    def test_native_editable_blocks_preserve_literal_text_lines_tables_and_embedded_images(self):
        html = f'''<div class="toolbar">PRIVATE TOOLBAR</div><section class="page">
        <div class="masthead"><img src="{URI}"><div><strong>PETROL ZONE</strong><span>SALES CONTRACT</span></div><div class="reference">SC-1<br>2026-10-19</div></div>
        <h1>Contract</h1><table><thead><tr><th style="width:40%">Product</th><th>Amount</th></tr></thead>
        <tbody><tr><td>Item<div class="product-description">Size &lt; 10 &amp; &gt; 2\n&lt;script&gt;literal&lt;/script&gt;\nLiteral &amp;amp; code</div></td><td>100</td></tr></tbody></table>
        </section><section class="page terms"><h2>Terms</h2><p>Clause &amp; text\nSecond line</p>
        <img class="seller-stamp" src="{URI}"></section>'''
        data = contract_docx(html)
        self.assertTrue(data.startswith(b'PK'))
        document = Document(BytesIO(data))
        self.assertGreaterEqual(len(document.tables), 2)
        self.assertIn('Item', document.tables[1].cell(1, 0).text)
        self.assertIn('Size < 10 & > 2\n<script>literal</script>\nLiteral &amp; code', document.tables[1].cell(1, 0).text)
        self.assertIn('Clause & text\nSecond line', [p.text for p in document.paragraphs])
        with ZipFile(BytesIO(data)) as archive:
            xml = archive.read('word/document.xml')
            tree = etree.fromstring(xml)
            self.assertGreater(len(tree.xpath('//w:t', namespaces=NS)), 10)
            self.assertTrue(tree.xpath('//w:tblHeader', namespaces=NS))
            self.assertTrue(tree.xpath('//w:pageBreakBefore', namespaces=NS))
            self.assertFalse(tree.xpath('//w:br[@w:type="page"]', namespaces=NS))
            self.assertNotIn(b'PRIVATE TOOLBAR', xml)
            self.assertNotIn(b'altChunk', xml)
            self.assertNotIn(b'documentProtection', archive.read('word/settings.xml'))
            media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
            self.assertIn(PNG, media)
            self.assertFalse(any('vba' in name.lower() for name in archive.namelist()))

    def test_never_reads_remote_or_local_images_or_builds_active_links(self):
        html = '''<section class="page"><p>&lt;img src="file:///private/secret"&gt; &lt;script&gt;x&lt;/script&gt;</p>
        <p><a href="https://example.invalid">link text</a></p><img src="https://example.invalid/stamp.png"><img src="file:///private/stamp.png"></section>'''
        data = contract_docx(html, blank=True)
        with ZipFile(BytesIO(data)) as archive:
            self.assertFalse(any(name.startswith('word/media/') for name in archive.namelist()))
            self.assertNotIn(b'TargetMode="External"', archive.read('word/_rels/document.xml.rels'))
            self.assertNotIn(b'hyperlink', archive.read('word/document.xml'))
        self.assertEqual(Document(BytesIO(data)).core_properties.title, 'Petrol Zone Sales Contract — Blank Template')

    def test_blank_writing_lines_and_control_characters(self):
        data = contract_docx('<section class="page"><p>Reference: <span class="line"></span></p><div class="ruled"></div><p>Safe\x00 &amp; text\tcontinued\nnext</p></section>')
        text = '\n'.join(p.text for p in Document(BytesIO(data)).paragraphs)
        self.assertIn('Reference: ____________________', text)
        self.assertIn('Safe & text\tcontinued\nnext', text)

    def test_saved_running_draft_or_cancelled_markers_cover_overflow_pages(self):
        for mark in ('DRAFT', 'CANCELLED'):
            data = contract_docx(f'<div id="header-html"><div class="pz-running">PETROL ZONE · SC-1 · {mark}</div></div><div id="footer-html"><div class="pz-footer">Saved Buyer · USD · SC-1</div></div><section class="page"><p>Content</p></section>')
            document = Document(BytesIO(data))
            self.assertIn(mark, document.sections[0].header.paragraphs[0].text)
            self.assertIn('Saved Buyer · USD · SC-1', document.sections[0].footer.paragraphs[0].text)
